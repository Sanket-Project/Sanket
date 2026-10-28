import json
from openpyxl import load_workbook

path = r"C:\Users\admin\Downloads\Redwine_Label_Fashion_Sales_and_Inventory_Record (1).xlsx"

_ALIASES = {
    "sku": {"sku","skucode","skuid","item","itemcode","itemsku","productcode","productsku","variantsku","code","externalid"},
    "quantity": {"quantity","qty","units","unitssold","unitvelocity","unitsvelocity","sold","count","volume","qtyordered","ordered","quantitysold","quantityordered","orderqty","orderedqty","salesqty","itemssold","totalunits","numberofunits","noofunits","unitsordered","piecesordered","pieces","pcs"},
    "price": {"sellingprice","price","unitprice","saleprice","amount","rate","retailprice","costprice","unitcost","cost"},
    "revenue": {"revenue","grossrevenue","totalrevenue","sales","total","linetotal","totalsales","saleamount","linerevenueamount","totalamount"},
    "cogs": {"cogs","totalcogs","costofgoodssold","costofsales"},
    "discount": {"discount","markdown","discountamount"},
    "timestamp": {"timestamp","date","saledate","orderdate","ordertimestamp","transactiondate","transactiontimestamp","datetime","time","soldat","createdat","invoicedate","salestimestamp"},
    "channel": {"channel","source","saleschannel","store","platform","checkoutchannel","multichannelcheckoutpath","ordertype"},
    "order_id": {"orderid","order","ordernumber","transactionid","orderno","invoice","invoiceid","invoicenumber","receiptid"},
    "region": {"region","country","state","market"},
    "available": {"availablestock","available","onhand","stock","stockonhand","currentstock","currentstockonhand","onhandunits","inventoryonhand","unitsonhand","physicalstock","physicalinventory","netstock"},
    "reserved": {"reservedstock","reserved","allocated","reservedunits"},
    "inbound": {"inboundstock","inbound","incoming","onorder","inboundunits"},
    "reorder_point": {"reorderpoint","reorderthreshold","safetystock","safetylevel","safetythreshold","reorderqty","minstock","minimumstock"},
    "warehouse": {"warehouseid","warehouse","location","hub","hubname","hubregion","fulfillmenthub","fulfillmentcenter","dc","site","binlocation","storagelocation"},
    "name": {"name","productname","title","product","itemname","skudescription"},
    "brand": {"brand","vendor","manufacturer","supplier"},
    "category": {"category","producttype","type","department","productcategory","categoryindexing","categoryindex"},
    "size": {"size","sizevariant","sizecode","productsize"},
    "description": {"description","desc","variant","option"},
    "gtin": {"gtin","barcode","ean","upc"},
}

def _norm(header):
    return "".join(ch for ch in header.lower() if ch.isalnum())

_all_aliases = set()
for a in _ALIASES.values():
    _all_aliases.update(a)

def _build_header_map(headers):
    found = {}
    for raw in headers:
        n = _norm(raw)
        for field_name, aliases in _ALIASES.items():
            if n in aliases and field_name not in found:
                found[field_name] = raw
    return found

wb = load_workbook(path, read_only=True, data_only=True)
report = {}
for name in wb.sheetnames:
    ws = wb[name]
    try:
        first_row = next(ws.iter_rows(values_only=True))
    except StopIteration:
        first_row = None
    score = 0
    normed = []
    if first_row is not None:
        for h in first_row:
            if h is not None:
                n = _norm(str(h))
                normed.append((str(h), n, n in _all_aliases))
                if n in _all_aliases:
                    score += 1
    report[name] = {"score": score, "cells": normed}
wb.close()

best_name, best_score = None, -1
for name, info in report.items():
    if info["score"] > best_score:
        best_score = info["score"]
        best_name = name
if best_score == 0:
    best_name = wb.sheetnames[0]

print(json.dumps({"scores": {k: v["score"] for k,v in report.items()}, "selected_sheet": best_name, "best_score": best_score}, indent=2))
print("---HEADER MAP for selected sheet (raw headers)---")
headers = [h for h,n,m in report[best_name]["cells"]]
print(json.dumps(_build_header_map(headers), indent=2))
print("---Per-cell detail for Sales Ledger---")
print(json.dumps(report.get("Sales Ledger", {}), indent=2))

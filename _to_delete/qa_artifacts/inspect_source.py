import json
from openpyxl import load_workbook

path = r"C:\Users\admin\Downloads\Redwine_Label_Fashion_Sales_and_Inventory_Record (1).xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
report = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        report[name] = {"rows": 0, "headers": []}
        continue
    headers = [str(h) if h is not None else None for h in rows[0]]
    data_rows = [r for r in rows[1:] if any(c is not None and str(c).strip() != "" for c in r)]
    blank_rows = (len(rows) - 1) - len(data_rows)
    report[name] = {
        
        "headers": headers,
        "data_row_count": len(data_rows),
        "blank_rows_skipped": blank_rows,
        "total_rows_incl_header": len(rows),
    }
wb.close()
print(json.dumps(report, indent=2, default=str))
with open(r"C:\Users\admin\Desktop\Prisma\qa_artifacts\source_inspection.json", "w") as f:
    json.dump(report, f, indent=2, default=str)
print("WROTE_REPORT_OK")

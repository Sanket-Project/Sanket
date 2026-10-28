"""Local-only state transitions; never submits payments or sends messages."""
import os,json,io
os.environ['SANKET_AUDIT_VARIANT']='clean-run'
from run_audit import C,call,OUT
_,a=call('actions_login','POST','/auth/dev-login',json=dict(email='fashion.audit-clean-run@example.com',password='AuditFashion2026!',tenant_slug='swecha-audit-clean-run'))
C.headers['Authorization']='Bearer '+a['access_token']
_,ds=call('decisions_before_actions','GET','/decisions')
orders=[d for d in ds['decisions'] if d['kind']=='purchase_order']
if orders:
    path='/decisions/'+orders[0]['id']+'/accept'
    call('accept_sample_recommendation','POST',path,json={'note':'LOCAL SYNTHETIC AUDIT ONLY - no supplier order placed'})
    call('repeat_accept_conflict','POST',path,json={})
if len(ds['decisions'])>1:
    d=next(d for d in ds['decisions'] if d['id']!=orders[0]['id'])
    call('dismiss_sample_recommendation','POST','/decisions/'+d['id']+'/dismiss',json={'reason':'Synthetic audit only'})
call('roi_after_sample_acceptance','GET','/roi/summary')
_,alerts=call('alerts_before_actions','GET','/alerts')
if alerts:
    call('acknowledge_sample_alert','PATCH','/alerts/'+alerts[0]['id']+'/acknowledge',json={'resolution_note':'Synthetic audit acknowledgement'})
    call('resolve_sample_alert','PATCH','/alerts/'+alerts[0]['id']+'/resolve',json={'resolution_note':'Synthetic audit resolution'})
call('inventory_negative_rejected','PUT','/inventory/levels',json={'items':[{'sku_id':'x','on_hand_units':-1}]})
call('scenario_invalid_demand','POST','/scenarios/simulate',json={'demand_shock_pct':1000})
call('empty_upload','POST','/integrations/upload',data={'kind':'sales'},files={'file':('empty.csv',b'','text/csv')})
call('ml_anonymous_denied','POST','http://127.0.0.1:8001/forecast',headers={'Authorization':''},json={'tenant_id':a['tenant_id'],'industry':'fashion','horizon':8})
# Excel inventory round-trip preserves the same audited positions.
from openpyxl import Workbook,load_workbook
m=json.loads((OUT/'sample-manifest.json').read_text())
wb=Workbook();ws=wb.active;ws.append(['sku','on_hand','location'])
for i,s in enumerate(m):
    for location,quantity in [('Boduppal',[30,240,80,190,35,70][i]),('Alwal',12)]: ws.append([s['code'],quantity,location])
buf=io.BytesIO();wb.save(buf);(OUT/'synthetic-inventory.xlsx').write_bytes(buf.getvalue())
call('analyze_excel','POST','/integrations/upload/analyze',data={'kind':'inventory'},files={'file':('inventory.xlsx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
call('upload_excel_inventory','POST','/integrations/upload',data={'kind':'inventory'},files={'file':('inventory.xlsx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
call('stock_after_excel','GET','/inventory/levels')
book=load_workbook(OUT/'forecast.xlsx',read_only=True)
info={s.title:{'rows':s.max_row,'columns':s.max_column} for s in book.worksheets}
(OUT/'excel-export-validation.json').write_text(json.dumps(info,indent=2))
print('Excel export validated:',info)

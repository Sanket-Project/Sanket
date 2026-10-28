"""Build review-only CSV contracts and audit documentation. Never writes application data."""
from pathlib import Path
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter
from decimal import Decimal
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[2]
ROOT.mkdir(parents=True, exist_ok=True)

def write(name, text):
    (ROOT / name).write_text(text.strip() + '\n', encoding='utf-8')

def csvout(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)

# Every field is declared once: name, type, necessity, definition/validation, example.
S = {}
def schema(name, fields):
    S[name] = [tuple(line.split('|', 4)) for line in fields.strip().splitlines()]

schema('products', '''schema_version|text|R|Exactly retail_v2.0.products|retail_v2.0.products
product_code|text(64)|R|Stable tenant-unique product/style identity; never infer from name|DEMO-KUR-001
product_name|text(200)|R|Display name; never a relationship key|Demo cotton printed kurti
brand_code|text(64)|O|Approved tenant brand lookup; blank means unknown|DEMO-BRAND
category_code|text(64)|R|Approved leaf category; its parent supplies subcategory hierarchy|KURTI
style_number|text(64)|O|Manufacturer style; not globally unique|DEMO-ST-101
description|text|O|Plain text description|Cotton printed straight kurti
collection_code|text(64)|O|Tenant collection lookup|DEMO-FESTIVE-26
season_code|text(64)|O|Tenant season lookup; not an inferred date|AW26
status|enum|R|active; seasonal; clearance; pre_launch; discontinued|active
hsn_code|text(8)|O|Verified HSN as text; required before taxable posting if applicable; blank in samples|
tax_class_code|text(64)|O|Effective-dated tax class lookup; required before taxable posting|DEMO-TAX-5
manufacturer_supplier_code|text(64)|O|Supplier FK for manufacturer only; purchasing suppliers use supplier_skus|DEMO-SUP-001
image_urls|json array|O|HTTPS URLs; no image fetched during import; blank unknown|
attributes_json|json object|O|Category-schema validated product defaults; SKU may override declared variant dimensions|{"fabric":"cotton","pattern":"floral","fit":"straight","sleeve":"three_quarter","neckline":"round","length":"knee","occasion":["casual"]}''')
schema('skus', '''schema_version|text|R|Exactly retail_v2.0.skus|retail_v2.0.skus
sku_code|text(64)|R|Tenant-unique stable sellable variant; preserve leading zeros|DEMO-KUR-001-BLU-M
product_code|text(64)|R|FK products.product_code; no catch-all parent in V2|DEMO-KUR-001
external_sku_id|text(128)|O|Source-specific identity; source_system is supplied by import profile|DEMO-EXT-001
color_code|text(64)|C|Required if category uses color dimension; approved lookup|BLUE
color_name|text(100)|O|Display label; must agree with color_code|Blue
size_code|text(32)|C|Required if category uses size; approved size-system lookup|M
size_system|text(32)|C|Required with size_code; label IN_ALPHA does not imply universal measurements|IN_ALPHA
barcode|text(64)|O|Internal barcode text; distinct namespace from GTIN|
gtin|text(14)|O|Only legitimate assigned 8/12/13/14 digit GTIN with valid check digit; blank sample|
mrp_incl_tax|decimal(18,2)|O|Nonnegative INR example; regulated maximum retail price; effective from price_valid_from|1575.00
selling_price_incl_tax|decimal(18,2)|O|Nonnegative default price; actual transaction price is separate|1050.00
reference_cost_ex_tax|decimal(18,4)|O|Optional planning/reference cost; never historical COGS|600.0000
currency|char(3)|R|ISO currency code; INR samples; never silently default USD|INR
price_valid_from|timestamptz|C|Required if MRP or selling price provided; creates price list record|2026-10-01T00:00:00+05:30
status|enum|R|active; inactive; discontinued|active
attributes_json|json object|O|Variant-only schema-validated attributes, e.g. measurements|{"bust_cm":96}''')
schema('inventory', '''schema_version|text|R|Exactly retail_v2.0.inventory; snapshot contract only|retail_v2.0.inventory
snapshot_id|text(128)|R|Source-stable snapshot identity; repeated keys are retry not new adjustment|DEMO-SNAP-001
as_of|timestamptz|R|Business measurement time with offset; not upload timestamp|2026-10-01T09:00:00+05:30
sku_code|text(64)|R|FK skus.sku_code|DEMO-KUR-001-BLU-M
location_code|text(64)|R|FK locations.location_code|demo-hyd-store
on_hand_qty|integer|R|Nonnegative physically held units including reserved/damaged/quarantine|20
reserved_qty|integer|R|Nonnegative subset of good on-hand, exclusive of damaged/quarantine|2
damaged_qty|integer|R|Nonnegative mutually exclusive unavailable stock bucket|1
quarantine_qty|integer|R|Nonnegative mutually exclusive quality-control stock bucket|0
in_transit_qty|integer|O|Destination-bound dispatched transfer units; reconcile transfer ledger; not on hand|0
open_po_qty|integer|O|Unreceived and uncancelled PO units; excludes transfer transit|0
source_system|text(64)|R|Registered external authority identifier; tenant scope server supplied|DEMO-POS
count_reference|text(128)|O|Approved physical count reference; mandatory for count adjustment posting|DEMO-COUNT-001
reported_stock_value_ex_tax|decimal(18,4)|O|Nonnegative source control value only; never silently overwrite cost ledger|12000.0000
currency|char(3)|C|Required when reported_stock_value_ex_tax present|INR''')
schema('sales', '''schema_version|text|R|Exactly retail_v2.0.sales|retail_v2.0.sales
source_system|text(64)|R|Registered source namespace; not a generic upload label|DEMO-POS
transaction_id|text(128)|R|Immutable source event identifier|DEMO-TXN-001
order_id|text(128)|R|Source order header identity|DEMO-ORDER-001
order_line_id|text(128)|R|Stable order line identity; cannot substitute SKU when repeated lines exist|DEMO-LINE-001
transaction_at|timestamptz|R|Recognition event time with offset; local business day from location timezone|2026-10-02T12:00:00+05:30
sku_code|text(64)|R|FK skus.sku_code; product derived through SKU|DEMO-KUR-001-BLU-M
location_code|text(64)|R|Selling location FK; fulfillment location optional separately|demo-hyd-store
fulfillment_location_code|text(64)|O|Stock issuing location FK; defaults to selling location only by confirmed profile|demo-hyd-store
channel_code|text(64)|R|Approved channel lookup; pos; web; marketplace examples|pos
quantity_sold|integer|R|Nonnegative; >0 for recognized sale; zero only for cancelled state|2
returns|integer|R|Exactly 0 here; physical returns posted only through returns.csv|0
unit_price_ex_tax|decimal(18,4)|R|Nonnegative selling price before discount excluding GST, per unit|1000.0000
mrp_incl_tax|decimal(18,2)|O|Historical MRP snapshot; not used to derive realized revenue|1575.00
discount_ex_tax|decimal(18,2)|R|Nonnegative total allocated LINE discount, not percentage and not per unit|200.00
tax_code|text(64)|R|Approved effective-dated classification; sample is illustrative only|DEMO-TAX-5
tax_rate_pct|decimal(7,4)|R|Nonnegative percentage; 5 means 5 percent; validate against configured tax code|5.0000
tax_amount|decimal(18,2)|R|Nonnegative GST total; inclusive/exclusive source conversion happens before this contract|90.00
currency|char(3)|R|ISO currency code; one currency per order|INR
order_status|enum|R|completed; cancelled; completed means recognized and fulfilled for this simple contract|completed
invoice_number|text(128)|O|Source tax invoice reference; required by configured fiscal posting policy|DEMO-INV-001
place_of_supply_state_code|text(2)|C|Required for GST posting; official state code text|36
payment_method|text(64)|O|Convenience only for exactly one tender; mixed requires payment child import|upi
payment_reference|text(128)|C|Required to create a payment; blank method alone does not imply payment|DEMO-PAY-001
payment_amount|decimal(18,2)|C|Nonnegative; repeat order payment reference must carry identical header data|1890.00''')
schema('returns', '''schema_version|text|R|Exactly retail_v2.0.returns|retail_v2.0.returns
source_system|text(64)|R|Registered source namespace|DEMO-POS
return_id|text(128)|R|Stable return header identity|DEMO-RET-001
return_line_id|text(128)|R|Stable return line identity|DEMO-RLINE-001
return_at|timestamptz|R|Approved return/credit recognition time; received disposition applied at this time in simple contract|2026-10-03T13:00:00+05:30
original_source_system|text(64)|C|Original namespace; mandatory unless approved unlinked return workflow|DEMO-POS
original_transaction_id|text(128)|C|Original sale event; mandatory for linked return|DEMO-TXN-001
original_order_line_id|text(128)|C|Original line; cumulative approved returns cannot exceed recognized sold units|DEMO-LINE-001
sku_code|text(64)|R|FK SKU; must match original line|DEMO-KUR-001-BLU-M
location_code|text(64)|R|Receiving location FK; may differ from original sale|demo-hyd-store
returns|integer|R|Strictly positive physical return quantity; never negative sales|1
return_value_ex_tax|decimal(18,2)|R|Nonnegative credit after original discount allocation, excluding GST|900.00
tax_code|text(64)|R|Original applicable tax classification; do not use today's rate|DEMO-TAX-5
tax_amount|decimal(18,2)|R|Nonnegative credited GST; separate fiscal eligibility validation|45.00
currency|char(3)|R|Must match original sale or use explicit FX workflow|INR
reason_code|text(64)|R|Approved reason e.g. size_issue; damaged; unwanted|size_issue
disposition|enum|R|restock; quarantine; damaged; no_receipt; simple file records physical outcome|restock
credit_note_number|text(128)|O|Fiscal/commercial credit reference; required for fiscal posting policy|DEMO-CN-001
refund_reference|text(128)|O|Optional one refund header; repeated ref must agree and allocated child lines deduplicate|DEMO-REF-001
refund_amount|decimal(18,2)|C|Nonnegative cash credit; required with refund reference; not additional revenue reversal|945.00
refund_method|text(64)|C|Required with refund reference; tender lookup|upi
refund_at|timestamptz|C|Required with refund reference; can differ from return time|2026-10-03T14:00:00+05:30
exchange_order_id|text(128)|O|Optional replacement sale order identity; both legs remain independent|
status|enum|R|approved; draft; rejected; only approved posts credit and receipt|approved''')
schema('purchases', '''schema_version|text|R|Exactly retail_v2.0.purchases; one row per document line event|retail_v2.0.purchases
source_system|text(64)|R|Registered source namespace|DEMO-ERP
document_type|enum|R|purchase_order; goods_receipt; supplier_invoice; purchase_return|purchase_order
document_id|text(128)|R|Stable identity within source and document type|DEMO-PO-001
document_line_id|text(128)|R|Stable line identity within document|DEMO-POL-001
document_at|timestamptz|R|Business posting/order time with offset|2026-10-04T10:00:00+05:30
supplier_code|text(64)|R|FK suppliers.supplier_code|DEMO-SUP-001
sku_code|text(64)|R|FK SKU; invoice summaries without SKU use separate summary lane|DEMO-KUR-001-BLU-M
location_code|text(64)|R|Destination/return-issuing location FK|demo-hyd-store
po_id|text(128)|C|PO header identity; required for PO and linked documents|DEMO-PO-001
po_line_id|text(128)|C|PO line identity; required for linked receipt/invoice/return|DEMO-POL-001
related_document_id|text(128)|C|For purchase_return: original receipt document ID; optional invoice/receipt matching elsewhere|
related_document_line_id|text(128)|C|Required with related document ID; original source is same namespace in simple contract|
quantity|integer|R|Strictly positive ordered/received/invoiced/returned units according to document_type|10
unit_cost_ex_tax|decimal(18,4)|R|Nonnegative unit base cost; realized receipt cost and invoice variance kept separately|600.0000
discount_ex_tax|decimal(18,2)|R|Nonnegative total line discount|0.00
freight_ex_tax|decimal(18,2)|R|Nonnegative allocated line freight; never repeat whole-header freight on each line|0.00
tax_code|text(64)|R|Configured purchase tax classification; sample illustrative|DEMO-TAX-5
tax_rate_pct|decimal(7,4)|R|Percentage; effective-date validate|5.0000
tax_amount|decimal(18,2)|R|Nonnegative total line tax after applicable allocation|300.00
currency|char(3)|R|ISO currency code; one currency per document|INR
invoice_number|text(128)|C|Required on supplier_invoice; unique per supplier/fiscal scope|
expected_delivery_date|date|O|ISO YYYY-MM-DD; useful on PO|2026-10-10
status|enum|R|draft; posted; cancelled; posted events immutable|posted''')
schema('suppliers', '''schema_version|text|R|Exactly retail_v2.0.suppliers|retail_v2.0.suppliers
supplier_code|text(64)|R|Tenant-unique stable identity|DEMO-SUP-001
supplier_name|text(200)|R|Synthetic sample name explicitly marked|DEMO ONLY Apparel Supplier
supplier_type|enum|R|manufacturer; wholesaler; distributor; other|manufacturer
gstin|text(15)|O|Verified tax registration; blank sample; syntax validity does not establish registration|
state_code|text(2)|C|Required for GST posting; state lookup|36
country_code|char(2)|R|ISO country code|IN
currency|char(3)|R|Default purchasing currency; transaction snapshots override explicitly|INR
payment_terms_days|integer|O|Nonnegative contractual due-day offset|30
default_lead_time_days|integer|O|Nonnegative fallback only; supplier_skus owns item-specific values|7
contact_email|text(254)|O|Optional email; no message sent by importing|demo-supplier@example.invalid
phone|text(32)|O|Optional contact number preserved as text|
address|text|O|Supplier address; blank synthetic sample|
status|enum|R|active; inactive; blocked|active''')
schema('locations', '''schema_version|text|R|Exactly retail_v2.0.locations|retail_v2.0.locations
location_code|text(64)|R|Tenant-unique stable identity; explicit aliases map legacy case variants|demo-hyd-store
location_name|text(200)|R|Human label|DEMO ONLY Hyderabad Store
location_type|enum|R|store; warehouse; virtual|store
parent_location_code|text(64)|O|FK location; no self-link or cycles|
timezone|text(64)|R|IANA timezone; business day calculated here|Asia/Kolkata
state_code|text(2)|C|Required for GST posting|36
country_code|char(2)|R|ISO country code|IN
address|text|O|Postal address; blank synthetic sample|
fulfills_orders|boolean|R|true or false only|true
status|enum|R|active; inactive|active''')

dictionary = []
for name, fields in S.items():
    for pos, (field, typ, required, rule, sample) in enumerate(fields, 1):
        dictionary.append([name, pos, field, typ, required, rule, sample])
csvout(ROOT / 'data_dictionary.csv', ['template','position','column','data_type','requirement','meaning_and_validation','synthetic_example'], dictionary)

samples = {name: [dict((f[0], f[4]) for f in fields)] for name, fields in S.items()}
samples['products'].append({**samples['products'][0], 'product_code':'DEMO-DUP-001','product_name':'Demo solid cotton dupatta','category_code':'DUPATTA','style_number':'DEMO-ST-102','attributes_json':'{"fabric":"cotton","pattern":"solid","length_cm":220}'})
samples['skus'].append({**samples['skus'][0], 'sku_code':'DEMO-DUP-001-NAT-OS','product_code':'DEMO-DUP-001','external_sku_id':'DEMO-EXT-002','color_code':'NATURAL','color_name':'Natural','size_code':'OS','size_system':'ONE_SIZE','mrp_incl_tax':'525.00','selling_price_incl_tax':'420.00','reference_cost_ex_tax':'200.0000','attributes_json':'{}'})
samples['inventory'].append({**samples['inventory'][0], 'sku_code':'DEMO-DUP-001-NAT-OS','on_hand_qty':'10','reserved_qty':'0','damaged_qty':'0','reported_stock_value_ex_tax':'2000.0000'})
base = samples['purchases'][0]
samples['purchases'] += [
    {**base,'document_type':'goods_receipt','document_id':'DEMO-GRN-001','document_line_id':'DEMO-GRNL-001','document_at':'2026-10-10T11:00:00+05:30'},
    {**base,'document_type':'supplier_invoice','document_id':'DEMO-PI-001','document_line_id':'DEMO-PIL-001','document_at':'2026-10-10T12:00:00+05:30','invoice_number':'DEMO-SUPINV-001','related_document_id':'DEMO-GRN-001','related_document_line_id':'DEMO-GRNL-001'},
    {**base,'document_type':'purchase_return','document_id':'DEMO-PR-001','document_line_id':'DEMO-PRL-001','document_at':'2026-10-11T12:00:00+05:30','quantity':'1','tax_amount':'30.00','related_document_id':'DEMO-GRN-001','related_document_line_id':'DEMO-GRNL-001'}]
for name, fields in S.items():
    headers = [f[0] for f in fields]
    csvout(ROOT / 'templates' / f'{name}_v2.csv', headers, [])
    csvout(ROOT / 'samples' / f'{name}_v2_SAMPLE_ONLY.csv', headers, [[row.get(h,'') for h in headers] for row in samples[name]])

md = ['# V2 CSV data dictionary', '', 'Proposal dated 28 October 2026. R = required; O = optional; C = conditionally required. Derived fields are excluded from write contracts. Each table fixes header order. Blank is unknown, never automatically zero. These are proposed contracts, not supported endpoints in the current application.', '', 'CSV uses UTF-8, RFC-style quoting, comma delimiter, decimal dot and no currency/grouping symbols. Identifiers are text. Piece quantities are integers. Financial totals round to 2 decimals using decimal half-up after line discount allocation. Unit costs/prices retain 4 decimals. Timestamps require ISO 8601 offset. A tenant is obtained from authenticated context, never an uploaded tenant_id.', '']
for name, fields in S.items():
    md += [f'## {name}', '', '```csv', ','.join(f[0] for f in fields), '```', '', '| Column | Type | Required | Meaning and validation |', '|---|---|---|---|']
    md += ['| ' + ' | '.join(f[:4]) + ' |' for f in fields]
    md += ['']
write('DATA_DICTIONARY.md', '\n'.join(md))

# Source manifest deliberately records headers and diagnostics without copying client identities.
inputs = [
    (Path('C:/Users/admin/Downloads/stock balence Report 2026.csv'), 'client report; origin inferred from existing audit; identity not independently authenticated'),
    (Path('C:/Users/admin/Downloads/Sales Data 2026.xlsx'), 'client report'),
    (Path('C:/Users/admin/Downloads/Gross Profit Data 2026.xlsx'), 'client report'),
    (Path('C:/Users/admin/Downloads/Purchasr Report 2026.xlsx'), 'client report'),
    (REPO / '_to_delete/qa_artifacts/qa_test_products_import.xlsx', 'synthetic QA fixture'),
    (REPO / '_to_delete/qa_artifacts/qa_test_inventory_import.xlsx', 'synthetic QA fixture'),
    (REPO / '_to_delete/qa_artifacts/Redwine_Label_Fashion_Sales_and_Inventory_Record (1).xlsx', 'archived fashion workbook; business authenticity unverified; excluded from live client controls'),
]
manifest = []
tables = []
for path, provenance in inputs:
    record = {'file':str(path), 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes':path.stat().st_size, 'provenance':provenance,'tables':[]}
    if path.suffix == '.csv':
        rows = list(csv.reader(path.open(encoding='utf-8-sig',newline='')))
        record['tables'].append({'sheet':'CSV','header_row':1,'headers':rows[0],'physical_rows':len(rows),'data_rows_before_total_filter':len(rows)-1})
    else:
        wb = load_workbook(path, read_only=True, data_only=False)
        for ws in wb.worksheets:
            rows = list(ws.values)
            hr = 9 if path.name in ('Sales Data 2026.xlsx','Gross Profit Data 2026.xlsx') else 2 if path.name=='Purchasr Report 2026.xlsx' else 1
            if ws.title=='Executive Summary':
                record['tables'].append({'sheet':ws.title,'physical_rows':ws.max_row,'physical_columns':ws.max_column,'note':'Presentation dashboard, not an import table; formula cache availability reviewed separately'})
                continue
            headers = [str(v) if v is not None else '' for v in rows[hr-1]]
            formulas = sum(isinstance(v,str) and v.startswith('=') for r in rows for v in r)
            cached = load_workbook(path,read_only=True,data_only=True)[ws.title]
            cached_rows = list(cached.values)
            missing_cache = sum(isinstance(v,str) and v.startswith('=') and cached_rows[i][j] is None for i,r in enumerate(rows) for j,v in enumerate(r))
            record['tables'].append({'sheet':ws.title,'header_row':hr,'headers':headers,'physical_rows':ws.max_row,'physical_columns':ws.max_column,'formula_cells':formulas,'formula_cells_without_cached_value':missing_cache})
    manifest.append(record)
    tables += [(path.name,t) for t in record['tables'] if 'headers' in t]
write('source_manifest.json', json.dumps(manifest,indent=2,ensure_ascii=False))

def number(v):
    if v is None or str(v).strip() == '': return None
    try: return Decimal(str(v))
    except Exception: return None

profiles = []
for path, provenance in inputs:
    if path.suffix == '.csv':
        allrows = list(csv.reader(path.open(encoding='utf-8-sig',newline='')))
        sections = [('CSV',allrows[0],allrows[1:-1],allrows[-1],(0,1))]
    else:
        wb = load_workbook(path,read_only=True,data_only=True)
        sections = []
        for ws in wb.worksheets:
            if ws.title=='Executive Summary': continue
            rr = list(ws.values)
            h = 9 if path.name in ('Sales Data 2026.xlsx','Gross Profit Data 2026.xlsx') else 2 if path.name=='Purchasr Report 2026.xlsx' else 1
            has_total = path.parent == Path('C:/Users/admin/Downloads')
            keycols = (0,1) if path.name in ('Sales Data 2026.xlsx','Gross Profit Data 2026.xlsx') else (0,2) if path.name=='Purchasr Report 2026.xlsx' else (0,5) if ws.title=='Warehouse Inventory' else (0,2) if ws.title=='Sales Ledger' else (0,)
            sections.append((ws.title,[str(v or '') for v in rr[h-1]],rr[h:-1] if has_total else rr[h:],rr[-1] if has_total else None,keycols))
    for sheet,headers,rows,total,keycols in sections:
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
        keys = Counter(tuple(str(r[i] or '') for i in keycols) for r in rows)
        pp = {'file':path.name,'sheet':sheet,'detail_rows':len(rows),'key_column_indices_0_based':keycols,'duplicate_candidate_key_groups':sum(n>1 for n in keys.values()),'duplicate_excess_rows':sum(n-1 for n in keys.values()),'columns':[]}
        for i,header in enumerate(headers):
            vals=[r[i] if i<len(r) else None for r in rows]
            nn=[number(v) for v in vals]
            numeric=[v for v in nn if v is not None]
            col={'header':header,'blank_count':sum(v is None or str(v).strip()=='' for v in vals),'nonblank_distinct_count':len({str(v) for v in vals if v is not None and str(v).strip()}),'numeric_count':len(numeric),'negative_numeric_count':sum(v<0 for v in numeric),'fractional_numeric_count':sum(v!=v.to_integral_value() for v in numeric)}
            # Only sum numeric measures, not style IDs or source item codes.
            if numeric and any(k in header.lower() for k in ('qty','quantity','amount','value','bal.val','profit','revenue','cost of goods')) and '%' not in header:
                col['detail_sum']=str(sum(numeric,Decimal(0)))
                control=number(total[i]) if total is not None and i<len(total) else None
                if control is not None:
                    col['footer_control']=str(control)
                    col['detail_minus_footer']=str(sum(numeric,Decimal(0))-control)
            pp['columns'].append(col)
        profiles.append(pp)
write('data_profile.json',json.dumps(profiles,indent=2))

product_headers = ['name','brand','category','subcategory','external_id','status','warehouse_inventory','lead_time','safety_stock','forecast_demand','monthly_sales']
sku_headers = ['product_external_id','sku_code','description','unit_cost','unit_price','currency','lead_time_days','lead_time','moq','safety_stock','reorder_point','gtin','external_id','warehouse_inventory','forecast_demand','monthly_sales']
inv_headers = ['sku_code','on_hand_units','inbound_units','reserved_units','location']
sales_headers = ['sku','name','quantity','timestamp','price','revenue','discount','channel','region','order_id','line_id','warehouse','returns','in_stock','stockout','currency']
tables = [('Application Products download',{'headers':product_headers}),('Application SKUs download',{'headers':sku_headers}),('Application Inventory download',{'headers':inv_headers}),('Application sales uploader supported fields (not a discovered CSV download)',{'headers':sales_headers})] + tables

def finding(field, table):
    f = field.lower().strip()
    rules = {
        'warehouse_inventory':('Unscoped stock quantity','Duplicates inventory ownership; lacks SKU/location/time','Remove from master; migrate only with confirmed grain and snapshot time','P0','Prevents inconsistent stock'),
        'forecast_demand':('Forecast scalar','No horizon, run, model, location or confidence','Move to forecast run/results; preserve legacy raw value only','P0','Forecast is not observed demand'),
        'monthly_sales':('Historical summary scalar','No month, gross/net definition or location','Import dated actual sales or retain labeled period aggregate','P0','Prevents fabricated history'),
        'lead_time':('Lead time','Ambiguous unit and duplicate lead_time_days','Map to supplier_skus.lead_time_days after confirmation','P1','Lead times differ by supplier and SKU'),
        'lead_time_days':('Replenishment lead time','SKU-global value misses supplier route','Use supplier SKU fallback plus lane lead times','P1','Correct supply planning'),
        'safety_stock':('Stock policy','Not physical stock; master-global policy','Move to SKU/location replenishment policy','P1','Store policies differ'),
        'safety/reorder threshold':('Combined policy threshold','Safety stock and reorder point conflated','Require separately defined safety_stock_qty and reorder_point_qty','P1','Avoid double counting buffers'),
        'reorder_point':('Reorder threshold','SKU-global policy','Move to SKU/location policy','P1','Correct multi-location planning'),
        'moq':('Minimum order quantity','No supplier or pack multiple','Move to supplier_skus with order_multiple','P1','Supplier-specific buying constraints'),
        'returns':('Returned units','Unlinked scalar lacks event date/value/refund','Nonnegative legacy only; V2 returns event with original line','P0','Avoid duplicate reversals'),
        'return qty':('Daily returned units','Store/day total cannot link original lines','Keep store/day summary; request line returns export','P0','No invented SKU returns'),
        'return amount':('Daily returned money','Tax and refund basis unclear','Preserve source summary and confirm basis','P0','Returns differ from cash refunds'),
        'on_hand_units':('Current physical stock','No mandatory snapshot time; inadequate status buckets','on_hand_qty plus as_of and location required','P0','Stops stale stock overwrites'),
        'inbound_units':('Incoming units','PO commitments and dispatched transfers conflated','Split open PO from transit; reconcile authoritative documents','P1','No duplicated supply'),
        'reserved_units':('Allocated physical units','No allocation lifecycle or oversubscription rule','Reservation records and disjoint stock state validation','P0','Prevents overselling'),
        'closing bal.qty':('Closing reported stock','Snapshot date absent; quantity basis needs confirmation','on_hand_qty only after confirmed semantics and as_of','P0','Never infer physical stock from ATS'),
        'closing bal.val':('Closing reported value','Cost method and tax basis not explicit','Store reported control value; reconcile valuation ledger','P0','No unsupported inventory valuation'),
        'quantity ordered':('Ordered units','Ordered is not fulfilled or recognized sale','Map only with recognition policy and status','P0','Avoid overstated sales'),
        'quantity':('Quantity','Meaning changes with document/report type','Require document_type or explicit sales quantity semantics','P0','Purchase returns cannot become sales'),
        'sales qty':('Reported sold units','Different file grains; cannot infer timestamps','Preserve report grain and period; request source transactions','P0','Correct demand eligibility'),
        'net qty':('Sold less returned units','Cannot reconstruct gross demand from net alone','Retain gross sales and returns separately','P0','Protects return rate and forecasts'),
        'stock no':('Source item code','Do not assume GTIN, style or globally unique identity','Tenant/source SKU crosswalk preserving exact text','P0','Reliable joins and leading zeros'),
        'style no':('Source style','May collide across brands/suppliers; not variant identity','Product/style attribute with explicit grouping crosswalk','P1','Avoid merging distinct products'),
        'product_external_id':('Parent product key','Frontend validation requires only SKU code','Require existing parent in V2; quarantine unresolved parents','P0','Prevents orphan or catch-all products'),
        'product id':('Parent source product','Repeated alongside SKU; may disagree','Validate against SKU parent; master-owned relationship','P0','Correct product aggregation'),
        'unit_cost':('Catalog reference unit cost','Not historical issue cost','reference_cost_ex_tax plus receipt valuation ledger','P0','Accurate historical COGS'),
        'unit product cost':('Source unit cost','Tax/cost method not declared','Preserve source and qualify cost basis; do not use current master','P0','Profit needs historical cost'),
        'cost price':('Source stock cost','Location variation and tax basis unknown','Preserve per-position reported cost; confirm tax and method','P0','No arbitrary catalog cost overwrites'),
        'unit_price':('Catalog price','MRP and selling price/tax basis not separated','Effective-dated selling price and MRP records','P1','Correct promotions and location pricing'),
        'retail price':('Reported retail price','Not proven MRP or actual sale price','Preserve source label; map only after confirmation','P1','Avoid mislabeled pricing'),
        'currency':('Monetary currency','SKU model defaults USD; uploads may omit currency','Require INR or explicit confirmed profile default','P0','No false currency conversion'),
        'gtin':('Trade item identifier','Text/check digit/allocation validation needed','Optional legitimate GS1 identifier; internal barcode separate','P1','SKU codes are not GTINs'),
        'size':('Variant size','No size system or measurement chart','size_code plus size_system and chart reference','P1','M is not universal sizing'),
        'company name':('Source branch label','May represent location rather than tenant/legal company','Confirmed location aliases within authenticated tenant','P0','Protects tenancy and location joins'),
        'voucher type':('Procurement document type','Signed returns require type-aware normalization','Map Purchase vs Purchase Return to explicit positive events','P0','Prevents double sign inversion'),
        'voucher number':('Voucher identifier','Must be scoped by source, supplier/type as appropriate','Immutable document identity and duplicate conflict checks','P0','Safe replays'),
        'party name':('Supplier label','Name not stable supplier identifier','Supplier crosswalk; preserve raw name and resolve ID','P1','No accidental supplier merges'),
        'gross profit % (based on cost)':('Markup percentage','Not margin; averaging invalid','Label markup; recompute from summed GP/cost','P0','Prevents overstated gross margin'),
        'gross profit % (based on sale)':('Sales margin percentage','Averaging row percentages invalid','Recompute summed GP/ex-tax revenue','P0','Weighted correct profitability'),
        'cost of goods (with tax)':('Tax-inclusive source cost','Recoverable tax may not be COGS','Retain reported source control; use validated ex-tax cost basis','P0','Consistent revenue/cost tax basis'),
        'cost of goods (without tax)':('Tax-exclusive source COGS','Period totals lack issue history','Import reported aggregate, not synthesized issue ledger','P0','Preserve auditability'),
        'total revenue':('Calculated revenue','Archived workbook formulas may have no cached values','Recalculate with validated engine or explicitly derive documented formula','P0','Blank cache is not zero revenue'),
        'total asset value (at cost)':('Calculated stock value','Formula cache and cost method need validation','Qualified source control only','P0','No false value from blank cache'),
        'price':('Sales unit selling price','Tax basis unspecified','Explicit unit_price_ex_tax; configured inclusive adapter','P0','Financial reconciliation'),
        'revenue':('Source gross line revenue','Gross/net/tax semantics not fixed','Explicit gross ex-tax control and recomputation','P0','No double discount'),
        'discount':('Sales discount','Amount vs percent and line vs unit ambiguous','Positive allocated line discount_ex_tax','P0','Correct net sales'),
        'timestamp':('Sales date/time','Timezone and recognition basis unspecified','Offset timestamp and location timezone','P0','Correct daily reporting'),
        'order_id':('Order identity','Not enough to identify repeated SKU lines','Require stable order_line_id and transaction event ID','P0','Safe deduplication'),
        'line_id':('Order line identity','Optional legacy identity permits ambiguous replay','Required immutable line ID in V2','P0','Corrections distinguishable from new events'),
    }
    if f in rules: return rules[f]
    if f in ('sku','sku_code'): return ('Sellable item identity','Must remain tenant-scoped text; unknown SKU policy differs by route','Resolve stable SKU FK; reject unknown V2 transaction references','P0','No invented catalog identity')
    if f in ('location','warehouse','warehouse code','fulfillment warehouse'): return ('Location reference','Free text aliases may duplicate a physical location','Canonical location FK with approved aliases','P0','Reliable stock and store metrics')
    if f in ('category','subcategory','sub category'): return ('Category classification','Free labels create taxonomy drift','Approved category hierarchy with raw-to-canonical mapping','P1','Reliable category reporting')
    if f in ('name','product name','item description','description'): return ('Item description','Must not generate identity or infer size/color without review','Keep display/raw description; explicit product and SKU codes','P1','Avoid variant merging')
    if f in ('brand','brand_code'): return ('Brand label','Free spelling variants fragment reporting','Tenant brand ID and confirmed alias mapping','P1','Reliable brand analysis')
    if f in ('external_id','status','channel','sales channel','region','in_stock','stockout','order date','voucher date','type','style number','fulfillment city'): return ('Source identity/status/dimension/date','Semantics require typed, scoped dictionary and null handling','Validate type and lookup; preserve raw source; avoid inferred facts','P1','Stable interoperable imports')
    if any(t in f for t in ('amount','value','cost','price','profit','revenue')): return ('Reported monetary measure','Tax, sign, grain or formula basis must be declared','Preserve reported control; canonical decimal and reconciliation policy','P0','Consistent financial basis')
    if f in ('card','cash','credit note','e-wallet','payment pendig'): return ('Daily tender/settlement summary','Not individual captured payments or refund allocations','Keep daily summary; request payment/refund references for reconciliation','P0','Avoid counting tenders as revenue')
    return ('Source report measure','Report-grain value is not an event identity or normalized master','Retain raw measure and declared grain; typed mapping before posting','P1','Avoid unsupported transaction reconstruction')

audit = ['# Critical audit of existing files', '', '28 October 2026. Static source/data audit; no live database was read or changed. Only the brief was attached in this chat. The four client reports below were found in Downloads and inspected directly. Three archived workbooks were inspected separately. Source hashes and exact headers are in source_manifest.json. No supplier identities, registrations or client sample rows are copied into the synthetic deliverables.', '', '## Decision', '', 'The current platform is an analytics and planning foundation, not yet a transaction-complete retail ERP. Preserve the existing mapping, replay identity, location and forecast capabilities. The highest risks are mixed report grains, ambiguous stock semantics, duplicate master quantities, incomplete historical costing and unlinked returns. A richer CSV alone will not fix them.', '', '## Verified application evidence', '', '- frontend/src/utils/csvImport.ts: Products requires name/category; SKUs requires sku_code only. Numeric master validation checks numeric conversion but not nonnegative/integer/domain constraints. Both master downloads include warehouse_inventory, forecast_demand and monthly_sales. SKU download includes both lead_time_days and lead_time.', '- frontend/src/api/inventory.ts: requires sku_code/on_hand_units, optionally inbound_units/reserved_units/location/source. The download omits as_of and source. Frontend checks nonnegative quantities; the server file mapper instead silently clamps negatives.', '- backend/app/services/integrations/file_import.py: map_sales_row rejects quantity < 0 and returns < 0, accepts zero demand, derives gross from price*quantity and net by subtracting discount. It does not define tax/return-value allocation here. Name-only rows can generate a SKU. map_inventory_row maps a broad available field into on_hand_units, clamps negative values, and defaults location. These are verified static findings, not a claim of a reproduced runtime failure.', '- backend/app/models/inventory.py: one current balance per tenant/SKU/location with as_of; available property clamps on_hand-reserved at zero. No disjoint damaged/quarantine buckets or movement ledger in this model.', '- backend/app/models/product.py: product attributes and SKU attributes already provide extension points. SKU has reference cost/price, defaults currency to USD, and SKU-global buying policy.', '- backend/app/models/location.py: locations and per-location forecasts already exist. V2 strengthens references and provenance rather than inventing these capabilities.', '- backend/app/services/sales_identity.py: transaction/period replay identities already exist. Fallback fingerprints and occurrence suffixes cannot establish real identity across edited overlapping files. Explicit source event IDs remain necessary.', '- backend/app/services/integrations/import_profiles.py and import_layout/import_review/import_dates/import_transforms: saved tenant mapping profiles, revision checks, file-bound confirmation and interpretation machinery already exist. Extend and harden these; do not replace them with an unrelated importer.', '', '## Observed files and reporting grains', '', '| Source | Actual grain and limitation | Safe use |', '|---|---|---|', '| stock balence Report 2026.csv | Company label / Stock No position. Stock date absent from columns. Description embeds some size/color tokens and ranges such as S TO XXL. | Confirm branch mapping, snapshot time and quantity basis. Do not expand assortment text into invented variants. |', '| Sales Data 2026.xlsx | Row 9 headers; store/day sales/return and tender totals; no SKU or order-line identifiers. | Daily store summary lane only. Never allocate these totals across SKUs. |', '| Gross Profit Data 2026.xlsx | Row 9 headers; SKU/store aggregate over title period; no individual sale dates. Includes both markup and margin. | Reported period profitability facts. Cannot create daily SKU demand. |', '| Purchasr Report 2026.xlsx | Row 2 headers; voucher/supplier quantities and amounts; no SKU lines or receiving locations. | Procurement summary lane. Cannot generate SKU receipts or PO fulfillment. |', '| qa_test_products_import.xlsx | Five synthetic master rows with planning quantities. | QA only; not client evidence. |', '| qa_test_inventory_import.xlsx | Four synthetic stock rows with no location/as_of. | Legacy QA only; cannot establish real opening stock. |', '| Redwine archived workbook | Sales Ledger and Warehouse Inventory plus Executive Summary. Formula cache gaps reported in manifest; authenticity unverified. | Structural test reference only; not proof of actual client activity. |', '', 'All raw workbook rows were scanned for structure/formulas. This audit does not assert independently re-executed reconciliation of every client total. Prior reconciliation results in docs/audit/2026-10-27-swecha-client-data/source-reconciliation.json remain prior evidence, not newly computed results.', '', '## Required and optional status today', '', 'Product download: name/category required, other columns optional. SKU download: sku_code required by frontend; parent optional there. Inventory: sku_code/on_hand_units required; location defaults in legacy flows. Sales universal import: sku/quantity/timestamp required at header validation, though mapper supports a name fallback. Other uploader fields are optional; this route mismatch needs a contract test. Reports are source exports, not application templates: requiredness for them is defined by a saved source profile, not guessed from populated cells.', '']
for label, table in tables:
    audit += [f'## {label}', '', '| Existing Field | Current Purpose | Problem | Recommended Change | Priority | Business Justification |', '|---|---|---|---|---|---|']
    for field in table['headers']:
        if not field: continue
        audit += ['| ' + ' | '.join([field, *finding(field,label)]) + ' |']
    audit += ['']
audit += ['## Gap assessment', '', '| Capability | Current evidence | V2 recommendation |', '|---|---|---|', '| Catalog | Products/SKUs and JSON attributes exist; parent optional in SKU upload | Enforce stable product parent and category-typed variants, including one default variant |', '| Inventory | Current balance with as_of; broad aliases and clamping | Explicit quantity semantics, immutable stock ledger, reservations and controlled snapshots |', '| Sales/returns | Gross units and returns fields, replay keys | Stable line/event identity, separate returns/credit/refund entities and recognition rules |', '| Procurement | Client voucher summaries; no procurement models found in inspected model directory | Supplier/PO/receipt/invoice/return lifecycle, retain aggregates separately |', '| Finance | Catalog cost and reported source GP | Historical valuation/COGS, tax basis and allocations |', '| Import interoperability | Mapping profiles and reviewed interpretation already implemented | Versioned canonical contracts, source authority, streaming and posting groups |', '| Forecasting | Run/result and location forecast structures exist | Feed only eligible dated actuals, preserve model/cutoff/horizon and censored demand |', '| Tenancy | Tenant IDs/RLS required by repository | Composite tenant FKs plus RLS and worker isolation |', '', 'Missing data is not zero. No report total, synthetic test row or forecast should silently become an observed SKU transaction.']
audit += ['', '## Full-row source profiling', '', 'data_profile.json records blank/distinct/numeric/negative/fractional counts for every inspected table column, candidate-key duplicates and independent detail sums against available footer controls. These diagnostics exclude confirmed title/header/footer rows. Candidate-key duplicates are review signals, not automatic proof of duplicated business events. Footer comparison is a separate fresh check from the prior client audit; it does not prove the business authenticity of either detail or footer.', '', '| File / sheet | Detail rows | Candidate duplicate groups | Excess rows |', '|---|---:|---:|---:|']
for pp in profiles:
    audit += [f"| {pp['file']} / {pp['sheet']} | {pp['detail_rows']} | {pp['duplicate_candidate_key_groups']} | {pp['duplicate_excess_rows']} |"]
audit += ['', 'Use the recorded detail-minus-footer differences to distinguish rounding and scope inconsistencies. Do not sum report percentages or price columns. Formula cells without cached values in the archived workbook are not zero-valued facts. The profile is read-only and does not repair or overwrite any source data.']
audit += ['', '## Specific data-quality observations', '', '- Stock: 16,519 detail positions and 38,419 closing units; no duplicate company/Stock No keys or negative numeric values in inspected detail. This does not establish the missing measurement date or prove that closing stock is physically sellable.', '- Sales: 121 store/day details; 2,683 sold units less 40 returned units = 2,643 net units. Sales amount INR 2,640,989 less returns INR 59,055 = INR 2,581,934 net. All these footer comparisons reconcile exactly. The already-net columns must not have returns subtracted again.', '- Gross profit: 15,524 details include 8 negative Sales Qty rows, 18 negative Net Sales Value rows, and 985 negative Gross Profit Amount rows. These are net period observations, not automatically invalid transactions. Preserve signed aggregate measures. Negative GP can represent genuine loss; never normalize every negative amount into a return. Gross sales demand cannot be recovered uniquely from these net quantities.', '- Purchase report: 266 detail vouchers include 29 negative quantities, 29 negative tax amounts and 29 negative net amounts. The source adapter must validate voucher type before converting return magnitudes to positive canonical purchase-return events. Negative purchase returns are not a reason to weaken the no-negative-sales rule.', '- Archived fashion workbook: 500 sale rows have 1,500 formula cells without cached results. Its 150 inventory rows have calculated asset-value/status fields with missing cached outputs. Numeric results must be explicitly calculated/verified or withheld; missing caches cannot become zero revenue, zero cost or zero inventory value.', '- No duplicate candidate keys were observed in these selected table grains. This is not evidence of globally unique transaction identity: the actual daily sales and period-profit reports do not contain transaction line keys.']
write('AUDIT.md','\n'.join(audit))

write('README.md', '''# Fashion retail V2 review package

Prepared 28 October 2026. This package is an architecture proposal, not an implemented application release. Existing V1 imports remain unchanged. V2 CSV files must not be uploaded to production until the V2 adapter is built and approved.

Read in this order:

1. AUDIT.md — verified current fields, per-field findings, source limitations and gap analysis.
2. ARCHITECTURE.md — ownership, normalized entities, stock/finance rules, ERD, trade-offs.
3. DATA_DICTIONARY.md or data_dictionary.csv — ordered complete field contracts for eight imports.
4. IMPORT_SPEC.md — validation, mapping, scale, replay, migration and acceptance examples.
5. ANALYTICS_AND_ROADMAP.md — KPI prerequisites, priorities, dependencies and acceptance criteria.
6. templates/ — eight header-only CSVs; samples/ — eight clearly synthetic CSVs.

All DEMO identities and example.invalid contacts are invented test values. GTIN, GSTIN and HSN fields are deliberately blank. DEMO-TAX-5 is an illustrative 5% configuration to exercise arithmetic, not a statement of the statutory rate for these goods. Configure verified effective-dated rates and HSN classifications before fiscal use. Taxable samples require that setup and verified HSN; they are not filing-ready invoices.

Import order: confirm tenant source profiles and brand/category/collection/season/tax/size/color/channel/reason lookups; suppliers; locations (parents first); products; SKUs; approved opening inventory; historical sales; returns; purchases in PO, receipt, invoice, return order. Opening inventory is as of 1 October before the sample sale. Opening stock represents prior history and is not a receipt from the sample PO. Historical imports at or before a later stock baseline do not mutate that baseline again. Procurement examples occur after the sale and return.

Sample expectations: sell two kurtis for INR 1,800 excluding tax after INR 200 line discount; GST INR 90; collect INR 1,890. Return one for INR 900 plus INR 45 credited tax; refund INR 945. Ex-tax net revenue INR 900, retained sale COGS INR 600, GP INR 300 and margin 33.3333% assuming confirmed historical issue cost INR 600/unit. PO 10, receipt 10, invoice 10, supplier return 1. PO remaining-to-receive = 0; net retained receipts = 9; supplier credit pending until separately documented. Final kurti physical stock = 20 - 2 + 1 + 10 - 1 = 28, including one damaged unit. With opening reservation 2 unchanged, ATS = 25. Dupatta has one one-size variant with 10 units and no sales.

source_manifest.json records inspected originals without copying client rows. SOURCES.md separates conventions from design choices. verification.json records package checks. The ZIP contains documentation, templates, dictionary, manifest and samples only; it excludes client files, backups, logs and this builder.
''')

write('SOURCES.md', '''# Sources and basis

Reviewed 28 October 2026. There is no universal retail CSV standard. Product/variant separation, explicit inventory states, distinct orders/payments/returns and historical costing are established patterns. The exact names, eight contracts, integer piece units, ledger posting policy and roadmap in this package are design recommendations for SANKET.

- [Shopify inventory states](https://help.shopify.com/en/manual/inventory-and-locations/fundamentals/inventory-states) and [developer inventory quantities](https://shopify.dev/docs/apps/build/inventory-management-apps/manage-quantities-states) support separating on-hand, unavailable, commitments and incoming. Our reserved definition and safety-stock policy are explicitly ours; they are not a claimed one-to-one Shopify mapping.
- [GS1 General Specifications, version 24 reference](https://ref.gs1.org/standards/genspecs/24.0.0/) supports distinct trade-item/color/size identifiers in apparel assortments. Consult the applicable current GS1 allocation rules when onboarding real barcodes. We do not fabricate assigned GTINs.
- [Microsoft Dynamics 365 inventory costing FAQ](https://learn.microsoft.com/en-us/dynamics365/supply-chain/cost-management/inventory-costing-faq) describes receipt/issue matching, valuation methods and cost adjustments. Moving weighted average is our recommended starting option, conditional on finance approval; the product does not need Dynamics-specific machinery.
- [CBIC Section 34](https://taxinformation.cbic.gov.in/content-page/explore-act/1000304/1000001) and [CBIC sectoral FAQs](https://cbic-gst.gov.in/sectoral-faq.html) support separate credit/debit-note treatment. A physical return does not itself prove eligibility to reduce GST. No sample GST rate or HSN classification is claimed as law. Fiscal policy must validate the applicable rules, dates and documentation.
- [PostgreSQL partitioning documentation](https://www.postgresql.org/docs/current/ddl-partitioning.html) explains partition-wide unique-key restrictions. The blanket comment in the existing sales model about partitioned tables/FKs should not be treated as a universal PostgreSQL limitation. Use supported tenant-aware keys and a separate unpartitioned identity registry where needed.

Local evidence: frontend/src/utils/csvImport.ts; frontend/src/api/inventory.ts; frontend/src/components/integrations/ImportUploader.tsx; backend/app/models/{product,inventory,sales,location}.py; backend/app/services/integrations/{file_import,import_profiles,import_layout,import_review,import_dates,import_transforms}.py; backend/app/services/sales_identity.py; prior client audit and source-reconciliation.json. Exact inspected data files, hashes, dimensions and header positions are in source_manifest.json. Local findings refer to the working tree, which already contained modifications; they are not an audit of a clean release commit.
''')

if __name__ == '__main__':
    print(json.dumps({'templates':len(S),'dictionary_fields':len(dictionary),'inspected_files':len(manifest)}))

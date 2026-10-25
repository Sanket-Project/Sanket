"""Every weekly number uses Monday 00:00 UTC weeks, whatever the session time zone (F02).

Each test pins the database session to India time on purpose: a sale at
20:00 UTC on Sunday 30 Aug is 01:30 on Monday 31 Aug in IST, and must still
belong to the week of 24 Aug.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime

from openpyxl import load_workbook
from sqlalchemy import text

from app.models.enums import IndustryCode
from app.services import data_quality_audit
from app.services.decision_builder import weekly_history

SUNDAY_NIGHT_UTC = datetime(2026, 8, 30, 20, 0, tzinfo=UTC)  # Monday 01:30 IST
NEXT_MONDAY_UTC = datetime(2026, 8, 31, 6, 0, tzinfo=UTC)


async def _sales(db_session, tenant, sku, rows):
    for when, units in rows:
        await db_session.execute(
            text(
                "INSERT INTO historical_sales (tenant_id, sku_id, industry, sale_time, channel, "
                "units_sold) VALUES (:t, :s, 'fashion', :w, 'pos', :u)"
            ),
            {"t": str(tenant), "s": str(sku), "w": when, "u": units},
        )
    await db_session.commit()


async def test_decision_history_buckets_in_utc(db_session, app_role_db, make_tenant):
    tenant, sku = await make_tenant(), uuid.uuid4()
    await _sales(db_session, tenant, sku, [(SUNDAY_NIGHT_UTC, 5), (NEXT_MONDAY_UTC, 3)])

    async with app_role_db.session(str(tenant)) as s:
        await s.execute(text("SET LOCAL TIME ZONE 'Asia/Kolkata'"))
        history = await weekly_history(s, tenant, IndustryCode.fashion, [sku])

    weekly, _ = history[sku]
    assert weekly == [5.0, 3.0]  # two UTC weeks, not one IST week of 8


async def test_data_quality_weeks_are_utc_weeks(db_session, app_role_db, make_tenant):
    tenant, sku = await make_tenant(), uuid.uuid4()
    await _sales(db_session, tenant, sku, [(SUNDAY_NIGHT_UTC, 5), (NEXT_MONDAY_UTC, 3)])

    async with app_role_db.session(str(tenant)) as s:
        await s.execute(text("SET LOCAL TIME ZONE 'Asia/Kolkata'"))
        report = await data_quality_audit.audit(s, tenant_id=tenant, industry="fashion")

    thin = next(f for f in report["findings"] if f["key"] == "thin_history")
    assert thin["sample"][0]["weeks_of_history"] == 2


async def test_hybrid_snapshot_records_its_calendar_and_cutoff(app_role_db, make_tenant):
    from app.schemas.trends import HybridForecastOut, HybridForecastSeriesOut
    from app.services.hybrid_forecast import _persist_accuracy_snapshot

    tenant = await make_tenant()
    out = HybridForecastOut.model_construct(
        industry=IndustryCode.fashion,
        horizon_weeks=2,
        generated_at=datetime(2026, 8, 30, 12, tzinfo=UTC),
        series=[
            HybridForecastSeriesOut.model_construct(
                sku_id=str(uuid.uuid4()), ds=["2026-08-31", "2026-09-07"],
                p10=[1.0, 1.0], p50=[2.0, 2.0], p90=[3.0, 3.0],
            )
        ],
        calendar="W-MON-UTC",
        history_end=date(2026, 8, 31),
    )

    await _persist_accuracy_snapshot(db=app_role_db, tenant_id=tenant, industry_code="fashion", out=out)

    async with app_role_db.session(str(tenant)) as s:
        metrics = await s.scalar(text("SELECT metrics FROM forecast_runs WHERE tenant_id = :t"),
                                 {"t": str(tenant)})
    assert metrics["calendar"] == "W-MON-UTC"
    assert metrics["history_end"] == "2026-08-31"
    assert metrics["mode"] == "hybrid"


async def test_forecast_export_names_the_week_convention(client, db_session, signed_in_workspace):
    tenant, headers = await signed_in_workspace()
    run_id, sku = uuid.uuid4(), uuid.uuid4()
    await db_session.execute(
        text(
            "INSERT INTO forecast_runs (id, tenant_id, industry, model_stack, horizon_weeks, "
            "granularity, status, completed_at, metrics) VALUES (:r, :t, 'fashion', "
            "ARRAY['chronos_zero_shot'], 1, 'weekly', 'completed', now(), "
            "'{\"calendar\": \"W-MON-UTC\"}')"
        ),
        {"r": str(run_id), "t": str(tenant)},
    )
    await db_session.execute(
        text(
            "INSERT INTO forecast_results (tenant_id, run_id, sku_id, forecast_date, p10, p50, p90, "
            "model_name) VALUES (:t, :r, :s, '2026-08-31', 1, 2, 3, 'chronos_zero_shot')"
        ),
        {"t": str(tenant), "r": str(run_id), "s": str(sku)},
    )
    await db_session.commit()

    r = await client.get("/api/v1/export/forecast.xlsx", headers=headers)

    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert "Week start (Mon, UTC)" in header
    assert "Week end (Sun)" in header
    row = {h: c.value for h, c in zip(header, ws[2], strict=False)}
    assert row["Week start (Mon, UTC)"] == "2026-08-31"
    assert row["Week end (Sun)"] == "2026-09-06"

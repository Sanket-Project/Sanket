"""Stock cover from demand, lead time and freshness — not a fixed 20-unit rule (F09).

Coverage never reads recommendations, so it is the same before and after a
decision refresh. SKU-wide cover pools every location; cover for one location
is only given when sales record that location.
"""

from __future__ import annotations

import json
import uuid
from datetime import UTC, date, datetime, timedelta

import pytest
from sqlalchemy import text

from app.services.coverage import coverage
from app.services.decision_builder import latest_completed_run

NOW = datetime(2026, 10, 26, 9, 0, tzinfo=UTC)  # Monday; the last complete week began 19 Oct
TRAILING_WEEKS = [date(2026, 8, 31) + timedelta(weeks=i) for i in range(8)]


async def _sku(db, tenant, code="KURTI-M", lead=14):
    pid, sid = uuid.uuid4(), uuid.uuid4()
    await db.execute(
        text("INSERT INTO products (id, tenant_id, industry, name, category, status) "
             "VALUES (:p, :t, 'fashion', 'Kurti', 'Kurtis', 'active')"),
        {"p": str(pid), "t": str(tenant)},
    )
    await db.execute(
        text("INSERT INTO skus (id, tenant_id, product_id, industry, sku_code, lead_time_days) "
             "VALUES (:s, :t, :p, 'fashion', :c, :l)"),
        {"s": str(sid), "t": str(tenant), "p": str(pid), "c": code, "l": lead},
    )
    return sid


async def _stock(db, tenant, sku, location, units, as_of=NOW - timedelta(days=1)):
    await db.execute(
        text("INSERT INTO inventory_levels (tenant_id, sku_id, industry, location, on_hand_units, "
             "reserved_units, inbound_units, source, as_of) VALUES (:t, :s, 'fashion', :l, :u, 0, 0, "
             "'synthetic_test', :a)"),
        {"t": str(tenant), "s": str(sku), "l": location, "u": units, "a": as_of},
    )


async def _weekly_sales(db, tenant, sku, per_week, weeks=TRAILING_WEEKS, location=None):
    for wk in weeks:
        await db.execute(
            text("INSERT INTO historical_sales (tenant_id, sku_id, industry, sale_time, channel, "
                 "units_sold, location) VALUES (:t, :s, 'fashion', :w, 'store', :u, :l)"),
            {"t": str(tenant), "s": str(sku), "u": per_week, "l": location,
             "w": datetime(wk.year, wk.month, wk.day, 12, tzinfo=UTC)},
        )


async def _forecast(db, tenant, sku, *, completed, weekly_p50, mode="zero_shot"):
    run = uuid.uuid4()
    await db.execute(
        text("INSERT INTO forecast_runs (id, tenant_id, industry, model_stack, horizon_weeks, "
             "granularity, status, completed_at, metrics) VALUES (:r, :t, 'fashion', "
             "ARRAY['chronos_zero_shot'], 8, 'weekly', 'completed', :c, CAST(:m AS jsonb))"),
        {"r": str(run), "t": str(tenant), "c": completed,
         "m": json.dumps({"mode": mode, "calendar": "W-MON-UTC", "history_end": "2026-10-26"})},
    )
    for i in range(8):
        await db.execute(
            text("INSERT INTO forecast_results (tenant_id, run_id, sku_id, forecast_date, p10, p50, "
                 "p90, model_name) VALUES (:t, :r, :s, :d, :p, :p, :p, 'chronos_zero_shot')"),
            {"t": str(tenant), "r": str(run), "s": str(sku), "p": weekly_p50,
             "d": date(2026, 10, 26) + timedelta(weeks=i)},
        )
    return run


async def _coverage(app_role_db, tenant):
    async with app_role_db.session(str(tenant)) as s:
        return await coverage(s, tenant_id=tenant, industry="fashion", now=NOW)


async def test_two_stores_with_different_stock_and_demand(db_session, app_role_db, make_tenant):
    tenant = await make_tenant()
    sku = await _sku(db_session, tenant)
    await _stock(db_session, tenant, sku, "Boduppal", 30)
    await _stock(db_session, tenant, sku, "Alwal", 190)
    await _weekly_sales(db_session, tenant, sku, 20, location="Boduppal")
    await _weekly_sales(db_session, tenant, sku, 5, location="Alwal")
    await db_session.commit()

    result = (await _coverage(app_role_db, tenant))["skus"][0]

    assert result["demand"]["basis"] == "trailing_actuals"
    assert result["demand"]["weekly_units"] == 25.0
    assert result["sku_wide"]["available_units"] == 220.0
    assert result["sku_wide"]["cover_days"] == 61.6
    assert result["sku_wide"]["status"] == "healthy"
    by_loc = {loc["location"]: loc for loc in result["locations"]}
    assert by_loc["Boduppal"]["cover_days"] == 10.5
    assert by_loc["Boduppal"]["status"] == "critical"
    assert by_loc["Alwal"]["cover_days"] == 266.0
    assert by_loc["Alwal"]["status"] == "healthy"


async def test_without_store_tagged_sales_only_pooled_cover_is_given(db_session, app_role_db, make_tenant):
    tenant = await make_tenant()
    sku = await _sku(db_session, tenant)
    await _stock(db_session, tenant, sku, "Boduppal", 30)
    await _stock(db_session, tenant, sku, "Alwal", 190)
    await _weekly_sales(db_session, tenant, sku, 25)  # no location recorded
    await db_session.commit()

    result = (await _coverage(app_role_db, tenant))["skus"][0]

    for loc in result["locations"]:
        assert loc["cover_days"] is None
        assert loc["status"] is None
        assert "all locations pooled" in loc["reason"]
    assert result["sku_wide"]["cover_days"] == 61.6


@pytest.mark.parametrize(("weeks", "basis"), [(8, "trailing_actuals"), (2, None)])
async def test_no_forecast_falls_back_to_actuals_only_with_enough_weeks(
    db_session, app_role_db, make_tenant, weeks, basis
):
    tenant = await make_tenant()
    sku = await _sku(db_session, tenant)
    await _stock(db_session, tenant, sku, "default", 40)
    await _weekly_sales(db_session, tenant, sku, 10, weeks=TRAILING_WEEKS[-weeks:])
    await db_session.commit()

    result = (await _coverage(app_role_db, tenant))["skus"][0]

    assert result["demand"]["basis"] == basis
    if basis is None:
        assert result["sku_wide"]["status"] == "unknown"
        assert result["sku_wide"]["cover_days"] is None


async def test_fresh_forecast_is_preferred_and_stale_one_ignored(db_session, app_role_db, make_tenant):
    fresh_tenant, stale_tenant = await make_tenant(), await make_tenant()
    for tenant, completed in ((fresh_tenant, NOW - timedelta(days=2)), (stale_tenant, NOW - timedelta(days=20))):
        sku = await _sku(db_session, tenant)
        await _stock(db_session, tenant, sku, "default", 140)
        await _weekly_sales(db_session, tenant, sku, 10)
        await _forecast(db_session, tenant, sku, completed=completed, weekly_p50=70)
    await db_session.commit()

    fresh = (await _coverage(app_role_db, fresh_tenant))["skus"][0]
    stale = (await _coverage(app_role_db, stale_tenant))["skus"][0]

    assert fresh["demand"]["basis"] == "forecast"
    assert fresh["demand"]["weekly_units"] == 70.0
    assert fresh["sku_wide"]["cover_days"] == 14.0
    assert fresh["sku_wide"]["status"] == "critical"
    assert stale["demand"]["basis"] == "trailing_actuals"


async def test_stale_stock_makes_cover_unknown(db_session, app_role_db, make_tenant):
    tenant = await make_tenant()
    sku = await _sku(db_session, tenant)
    await _stock(db_session, tenant, sku, "default", 500, as_of=NOW - timedelta(days=40))
    await _weekly_sales(db_session, tenant, sku, 10)
    await db_session.commit()

    result = (await _coverage(app_role_db, tenant))["skus"][0]

    assert result["sku_wide"]["status"] == "unknown"
    assert "40 days" in result["sku_wide"]["reason"]


async def test_a_backtest_is_never_the_latest_forecast(db_session, app_role_db, make_tenant):
    tenant = await make_tenant()
    sku = await _sku(db_session, tenant)
    real = await _forecast(db_session, tenant, sku, completed=NOW - timedelta(days=3), weekly_p50=70)
    await _forecast(db_session, tenant, sku, completed=NOW - timedelta(hours=1), weekly_p50=999,
                    mode="backtest")
    await _stock(db_session, tenant, sku, "default", 140)
    await db_session.commit()

    async with app_role_db.session(str(tenant)) as s:
        from app.models.enums import IndustryCode

        run = await latest_completed_run(s, tenant, IndustryCode.fashion)
    result = (await _coverage(app_role_db, tenant))["skus"][0]

    assert run.id == real
    assert result["demand"]["weekly_units"] == 70.0


async def test_coverage_is_tenant_scoped(db_session, app_role_db, make_tenant):
    a, b = await make_tenant(), await make_tenant()
    sku = await _sku(db_session, a)
    await _stock(db_session, a, sku, "default", 10)
    await db_session.commit()

    assert (await _coverage(app_role_db, b))["skus"] == []
    assert len((await _coverage(app_role_db, a))["skus"]) == 1


async def test_endpoint_is_independent_of_decision_refresh(client, signed_in_workspace, db_session):
    tenant, headers = await signed_in_workspace()
    sku = await _sku(db_session, tenant)
    await _stock(db_session, tenant, sku, "default", 30, as_of=datetime.now(UTC))
    await db_session.commit()

    before = await client.get("/api/v1/inventory/coverage", headers=headers)
    await client.post("/api/v1/decisions/refresh", headers=headers)
    after = await client.get("/api/v1/inventory/coverage", headers=headers)

    assert before.status_code == 200, before.text
    assert before.json()["skus"] == after.json()["skus"]
    assert before.json()["skus"][0]["sku_wide"]["status"] == "unknown"  # no demand basis yet


async def test_export_never_ships_a_backtest_as_the_forecast(client, signed_in_workspace, db_session):
    import io

    from openpyxl import load_workbook

    tenant, headers = await signed_in_workspace()
    sku = await _sku(db_session, tenant)
    await _forecast(db_session, tenant, sku, completed=datetime.now(UTC) - timedelta(days=1),
                    weekly_p50=70)
    await _forecast(db_session, tenant, sku, completed=datetime.now(UTC), weekly_p50=999,
                    mode="backtest")
    await db_session.commit()

    r = await client.get("/api/v1/export/forecast.xlsx", headers=headers)

    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    p50s = {row[header.index("P50 (Median)")] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert p50s == {70.0}

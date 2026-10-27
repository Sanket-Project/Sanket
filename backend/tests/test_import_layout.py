import io
import zipfile

import pytest
from openpyxl import Workbook


def workbook(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def discover(content, kind="sales", filename="file.xlsx", **kwargs):
    from app.services.integrations.import_layout import discover_tables
    return discover_tables(filename, content, kind, **kwargs)


def test_sheet_selection_depends_on_kind_not_first_sheet():
    raw = workbook({"Summary": [["Company report"], ["Prepared for review"]],
                    "Sales": [["SKU", "Qty", "Date"], ["A", 3, "2026-01-01"]],
                    "Stock": [["SKU", "Available Stock", "Warehouse"], ["A", 17, "ST1"]]})
    assert discover(raw)[0].sheet == "Sales"
    assert discover(raw, "inventory")[0].sheet == "Stock"
    assert discover(raw, "inventory")[0].records[0]["c2"] == 17


def test_inventory_sheet_ranking_with_unfamiliar_stock_header():
    raw = workbook({"Sales Ledger": [["SKU", "Qty", "Date", "Price"], ["A", 3, "2026-01-01", 10]],
                    "Warehouse Inventory": [["SKU", "Stock On Hand (Qty)", "Warehouse Code"], ["A", 17, "ST1"]]})
    assert discover(raw, "inventory")[0].sheet == "Warehouse Inventory"


@pytest.mark.parametrize("headers", [["SKU", "Qty", "Date"], ["उत्पाद", "मात्रा", "दिनांक"], ["商品", "数量", "日付"]])
def test_csv_delayed_multilingual_header_detection(headers):
    content = ("Monthly report\nNotes for review\n\n" + ",".join(headers) + "\nA,2,2026-01-01\nB,0,2026-01-02\n").encode()
    table = discover(content, filename="report.csv")[0]
    assert table.header_start == 4
    assert len(table.records) == 2
    assert table.row_numbers == [5, 6]


def test_duplicate_blank_headers_keep_physical_columns():
    table = discover(b"SKU,Qty,Date,Note,Note,\nA,2,2026-01-01,first,second,last\n", filename="a.csv")[0]
    assert len({c.id for c in table.columns}) == 6
    assert table.records[0]["c5"] == "second"
    assert table.records[0]["c6"] == "last"


def test_manual_header_beyond_search_and_invalid_selection():
    raw = workbook({"Data": [["Notes"]] * 110 + [["SKU", "Qty", "Date"], ["A", 2, "2026-01-01"]]})
    table = discover(raw, header_overrides={"Data": [111, 111]})[0]
    assert table.header_start == 111
    with pytest.raises(ValueError, match="header"):
        discover(raw, header_overrides={"Data": [999, 999]})


def test_merged_headers_and_data_are_distinguished_and_formulas_warn():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Report title"])
    ws.merge_cells("A1:D1")
    ws.append(["Identity", None, "Stock", None])
    ws.merge_cells("A2:B2")
    ws.merge_cells("C2:D2")
    ws.append(["SKU", "Name", "Available", "Reserved"])
    ws.append(["A", "Shirt", 12, "=1+1"])
    ws.append([None, "Coat", 3, 0])
    ws.merge_cells("A4:A5")
    stream = io.BytesIO()
    wb.save(stream)
    table = discover(stream.getvalue(), "inventory", header_overrides={"Stock": [2, 3]})[0]
    assert table.columns[0].label == "Identity / SKU"
    assert table.columns[1].label == "Identity / Name"
    assert table.records[1]["c1"] is None
    assert any("D4" in warning and "cached" in warning for warning in table.warnings)
    assert any("A5" in warning and "merged" in warning for warning in table.warnings)


def test_exclusions_do_not_drop_total_named_products_or_zero_sales():
    raw = b"SKU,Qty,Date\nA,2,2026-01-01\nSKU,Qty,Date\nSubtotal,2,\nTotal,0,2026-01-02\n,2,\n"
    table = discover(raw, filename="a.csv")[0]
    assert len(table.records) == 3
    assert table.records[1]["c1"] == "Total"
    assert [e["reason"] for e in table.exclusions] == ["repeated_header", "subtotal"]
    assert table.review_rows == [6]


def test_inventory_sku_named_total_is_not_silently_excluded():
    table = discover(b"SKU,Available Stock,Warehouse\nTotal,7,ST1\nA,3,ST1\n", "inventory", filename="inventory.csv")[0]
    assert len(table.records) == 2
    assert table.records[0]["c1"] == "Total"


def test_fingerprint_ignores_filename_offset_values_counts_but_preserves_shape():
    a = discover(b"SKU,Qty,Date\nA,2,2026-01-01\n", filename="a.csv")[0]
    b = discover(b"Title\nSKU,Qty,Date\nB,5,2026-02-01\nC,6,2026-03-01\n", filename="b.csv")[0]
    changed = discover(b"Qty,SKU,Date\n2,A,2026-01-01\n", filename="a.csv")[0]
    assert a.fingerprint == b.fingerprint
    assert a.fingerprint != changed.fingerprint


def test_limits_reject_instead_of_truncating():
    with pytest.raises(ValueError, match="columns"):
        discover((",".join(["SKU"] * 201) + "\n" + ",".join(["A"] * 201)).encode(), filename="a.csv")
    with pytest.raises(ValueError, match="sheets"):
        discover(workbook({f"s{i}": [["SKU", "Qty"], ["A", 2]] for i in range(21)}))
    with pytest.raises(ValueError, match="workbook"):
        discover(b"not a zip")


def test_field_registry_only_includes_persisted_fields():
    from app.services.integrations.import_layout import fields_for
    assert "warehouse" in fields_for("sales")
    assert "available" in fields_for("inventory")
    assert "reorder_point" not in fields_for("inventory")
    assert "cogs" not in fields_for("products")
    with pytest.raises(ValueError):
        fields_for("unknown")


def test_corrupt_sheet_xml_returns_clear_workbook_error():
    raw = workbook({"Data": [["SKU", "Qty", "Date"], ["A", 2, "2026-01-01"]]})
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(stream, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<worksheet>broken" if name == "xl/worksheets/sheet1.xml" else source.read(name))
    with pytest.raises(ValueError, match="workbook"):
        discover(stream.getvalue())


def test_header_search_materializes_only_one_full_table_per_sheet(monkeypatch):
    import app.services.integrations.import_layout as layout
    original = layout._candidate
    full_tables = []
    def tracked(*args, **kwargs):
        if not kwargs.get("preview_only", False):
            full_tables.append(True)
        return original(*args, **kwargs)
    monkeypatch.setattr(layout, "_candidate", tracked)
    raw = workbook({"Data": [["Company report", "Notes"] for _ in range(50)] + [["SKU", "Qty", "Date"], ["A", 2, "2026-01-01"]]})
    assert discover(raw)[0].header_start == 51
    assert len(full_tables) == 1


def test_numeric_header_suffixes_and_sku_codes_are_not_misclassified_as_values():
    table = discover(b"SKU,Qty 2026,Date\nA123,2,2026-01-01\n", filename="sales.csv")[0]
    assert table.header_start == 1
    assert table.columns[0].kind == "text"
